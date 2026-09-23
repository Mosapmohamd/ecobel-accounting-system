from datetime import datetime
from io import BytesIO
from typing import Optional, Literal

from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session, joinedload

from .. import models, auth
from ..database import get_db

router = APIRouter(prefix="/reports/export", tags=["Reports"], dependencies=[Depends(auth.get_current_user)])

FINANCE_TYPE_LABEL = {"income": "إيراد", "expense": "مصروف"}
SOURCE_LABEL = {"website": "الموقع", "b2b": "جملة B2B", "spending": "مصروفات عامة", None: "مصروفات عامة"}
STATUS_LABEL = {"pending": "قيد التجهيز", "shipped": "في الطريق", "delivered": "تم التوصيل", "cancelled": "ملغي"}


def _xlsx_response(wb: Workbook, filename: str) -> StreamingResponse:
    for ws in wb.worksheets:
        for col_cells in ws.columns:
            length = max((len(str(c.value)) if c.value is not None else 0) for c in col_cells)
            ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(max(length + 2, 10), 50)
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    # HTTP headers must be latin-1 — an Arabic filename has to go through
    # RFC 5987 encoding (filename*=UTF-8''...); a plain ASCII filename is
    # kept too as a fallback for older clients that don't parse filename*.
    from urllib.parse import quote
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=report.xlsx; filename*=UTF-8''{quote(filename)}"},
    )


@router.get("/finance")
def export_finance(
    source: Optional[Literal["website", "b2b", "spending"]] = None,
    type: Optional[models.FinanceEntryType] = None,
    date_from: Optional[datetime] = None,
    date_to: Optional[datetime] = None,
    db: Session = Depends(get_db),
):
    q = db.query(models.FinanceEntry)
    if type:
        q = q.filter(models.FinanceEntry.type == type)
    if source:
        if source == "spending":
            q = q.filter((models.FinanceEntry.source == "spending") | (models.FinanceEntry.source.is_(None)))
        else:
            q = q.filter(models.FinanceEntry.source == source)
    if date_from:
        q = q.filter(models.FinanceEntry.entry_date >= date_from)
    if date_to:
        q = q.filter(models.FinanceEntry.entry_date <= date_to)
    entries = q.order_by(models.FinanceEntry.entry_date.desc()).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "الحركات المالية"
    ws.append(["النوع", "المصدر", "الفئة", "المبلغ", "الوصف", "التاريخ"])
    for e in entries:
        ws.append([
            FINANCE_TYPE_LABEL.get(e.type.value, e.type.value),
            SOURCE_LABEL.get(e.source, e.source or "مصروفات عامة"),
            e.category,
            e.amount,
            e.description or "",
            e.entry_date.strftime("%Y-%m-%d %H:%M") if e.entry_date else "",
        ])
    return _xlsx_response(wb, "التقرير_المالي.xlsx")


@router.get("/online-orders")
def export_online_orders(
    status: Optional[models.OrderStatus] = None,
    db: Session = Depends(get_db),
):
    q = db.query(models.Order).options(joinedload(models.Order.items))
    if status:
        q = q.filter(models.Order.status == status)
    orders = q.order_by(models.Order.created_at.desc()).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "طلبات الموقع"
    ws.append([
        "رقم الطلب", "اسم العميل", "التليفون", "المدينة", "العنوان", "الحالة",
        "المنتجات", "الإجمالي الفرعي", "الخصم", "الشحن", "الإجمالي", "التاريخ",
    ])
    for o in orders:
        items_str = " | ".join(f"{it.product_name} × {it.quantity}" for it in o.items)
        ws.append([
            o.order_number, o.customer_name, o.customer_phone, o.city or "", o.shipping_address,
            STATUS_LABEL.get(o.status.value, o.status.value), items_str,
            o.subtotal, o.discount_amount, o.shipping_fee, o.total_amount,
            o.created_at.strftime("%Y-%m-%d %H:%M") if o.created_at else "",
        ])
    return _xlsx_response(wb, "طلبات_الموقع.xlsx")


@router.get("/b2b-orders")
def export_b2b_orders(db: Session = Depends(get_db)):
    orders = (
        db.query(models.B2BOrder)
        .options(joinedload(models.B2BOrder.items), joinedload(models.B2BOrder.customer))
        .order_by(models.B2BOrder.created_at.desc())
        .all()
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "أوردرات B2B"
    ws.append(["العميل", "المنتجات والكميات", "خصم إضافي لمرة واحدة", "الإجمالي بعد الخصم", "ملاحظات", "التاريخ"])
    for o in orders:
        items_str = " | ".join(f"{it.product_id} × {it.quantity}" for it in o.items)
        ws.append([
            o.customer.name if o.customer else "—",
            items_str,
            f"{o.extra_discount_percentage}%" if o.extra_discount_percentage else "—",
            o.total_amount,
            o.note or "",
            o.created_at.strftime("%Y-%m-%d %H:%M") if o.created_at else "",
        ])
    return _xlsx_response(wb, "أوردرات_B2B.xlsx")


@router.get("/inventory")
def export_inventory(db: Session = Depends(get_db)):
    products = db.query(models.Product).order_by(models.Product.name).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "المخزون"
    ws.append(["المنتج", "الفئة", "SKU", "سعر البيع", "الكمية", "حد إعادة الطلب", "الحالة", "نشط"])
    for p in products:
        status_ar = {"ok": "متوفر", "low": "منخفض", "out": "نفذ"}[p.stock_status]
        ws.append([
            p.name, p.category.name if p.category else "", p.sku or "",
            p.sale_price, p.quantity, p.low_stock_threshold, status_ar,
            "نعم" if p.is_active else "لا",
        ])
    return _xlsx_response(wb, "تقرير_المخزون.xlsx")
