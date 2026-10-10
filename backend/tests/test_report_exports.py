"""Excel exports: each report is its own workbook with a descriptive,
dated filename (<report>[_<filter>]_<YYYY-MM-DD>.xlsx); the workbook
contents are unchanged."""
import io
import os
import re
import sys
import tempfile
from datetime import date

if "app.database" not in sys.modules:
    os.environ["DATABASE_URL"] = f"sqlite:///{os.path.join(tempfile.mkdtemp(), 'test.db')}"
    os.environ.pop("SECRET_KEY", None)

import pytest  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from app import auth, cairo_time, models  # noqa: E402
from app.database import SessionLocal  # noqa: E402
from app.main import app  # noqa: E402
from app.routers.report_exports import export_filename  # noqa: E402

client = TestClient(app)
H = {"Authorization": f"Bearer {auth.create_access_token({'sub': 'exports_staff'})}"}
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# Each export: URL -> (report name in the filename, sheet title, header row).
# The titles/headers are what the reports produced before the filenames
# changed — they must stay exactly the same.
EXPORTS = {
    "/reports/export/finance": ("Finance_Report", "الحركات المالية",
                                ["النوع", "المصدر", "الفئة", "المبلغ", "الوصف", "التاريخ"]),
    "/reports/export/online-orders": ("Online_Orders_Report", "طلبات الموقع",
                                      ["رقم الطلب", "اسم العميل", "التليفون", "المدينة", "العنوان", "الحالة", "المنتجات",
                                       "الإجمالي الفرعي", "الخصم", "الشحن", "الإجمالي", "التاريخ"]),
    "/reports/export/b2b-orders": ("B2B_Orders_Report", "أوردرات B2B",
                                   ["العميل", "المنتجات والكميات", "خصم إضافي لمرة واحدة", "الإجمالي بعد الخصم", "ملاحظات", "التاريخ"]),
    "/reports/export/inventory": ("Inventory_Report", "المخزون",
                                  ["المنتج", "الفئة", "SKU", "سعر البيع", "الكمية", "حد إعادة الطلب", "الحالة", "نشط"]),
}
NAME = re.compile(r'attachment; filename="([A-Za-z0-9_]+_(\d{4}-\d{2}-\d{2})\.xlsx)"$')


@pytest.fixture(scope="module", autouse=True)
def seed():
    db = SessionLocal()
    db.add(models.User(username="exports_staff", hashed_password=auth.hash_password("x")))
    db.add(models.Category(id="xc", name="فئة التصدير"))
    db.add(models.Product(id="xp", name="منتج التصدير", category_id="xc", sale_price=100, quantity=5))
    db.add(models.FinanceEntry(type=models.FinanceEntryType.income, category="sales", source="website", amount=50, description="بيع"))
    db.commit()
    db.close()


def filename_of(res):
    m = NAME.match(res.headers["content-disposition"])
    assert m, res.headers["content-disposition"]
    return m.group(1), m.group(2)


@pytest.mark.parametrize("url", list(EXPORTS))
def test_each_report_is_its_own_workbook_with_unchanged_contents(url):
    _, title, header = EXPORTS[url]
    res = client.get(url, headers=H)
    assert res.status_code == 200 and res.headers["content-type"] == XLSX
    wb = load_workbook(io.BytesIO(res.content))
    assert [ws.title for ws in wb.worksheets] == [title]  # one report per file
    assert [c.value for c in next(wb.worksheets[0].iter_rows(max_row=1))] == header


@pytest.mark.parametrize("url", list(EXPORTS))
def test_filename_names_report_and_todays_date(url):
    report = EXPORTS[url][0]
    name, day = filename_of(client.get(url, headers=H))
    assert name == f"{report}_{day}.xlsx"
    assert day == cairo_time.cairo_today().isoformat()  # the Cairo date, on any server clock
    assert name != "report.xlsx" and "report.xlsx" not in name


def test_every_report_gets_a_different_filename():
    names = {filename_of(client.get(url, headers=H))[0] for url in EXPORTS}
    assert len(names) == len(EXPORTS)


@pytest.mark.parametrize("query,expected", [
    ("?source=website", "Finance_Report_Website_"),
    ("?source=b2b", "Finance_Report_B2B_"),
    ("?source=spending", "Finance_Report_Spending_"),
    ("?type=expense", "Finance_Report_Expense_"),
])
def test_filtered_finance_exports_say_which_filter(query, expected):
    assert filename_of(client.get("/reports/export/finance" + query, headers=H))[0].startswith(expected)


def test_online_orders_status_filter_in_filename():
    assert filename_of(client.get("/reports/export/online-orders?status=shipped", headers=H))[0].startswith(
        "Online_Orders_Report_Shipped_")


def test_date_is_the_export_day_and_filename_is_filesystem_safe():
    name = export_filename("Inventory_Report", today=date(2026, 1, 2))
    assert name == "Inventory_Report_2026-01-02.xlsx"
    assert re.fullmatch(r"[A-Za-z0-9_\-.]+", name)  # no spaces, slashes, colons or non-ASCII


def test_filename_date_is_cairo_date_not_server_utc_date(monkeypatch):
    """21:30 UTC on 10 Oct is already 00:30 on 11 Oct in Cairo (UTC+3)."""
    from datetime import datetime, timezone
    monkeypatch.setattr(cairo_time, "now_utc", lambda: datetime(2026, 10, 10, 21, 30, tzinfo=timezone.utc))
    name, day = filename_of(client.get("/reports/export/inventory", headers=H))
    assert name == "Inventory_Report_2026-10-11.xlsx" and day == "2026-10-11"


def test_exports_still_require_staff_login():
    assert client.get("/reports/export/inventory").status_code == 401
